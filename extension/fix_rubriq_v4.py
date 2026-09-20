# -*- coding: utf-8 -*-
"""Repair the meaning-changing errors introduced by the Rubriq edit of v4.

The edit is kept: US spelling suits Epilepsia Open, and most of the rephrasing
is an improvement. What is fixed here are places where the copyedit reversed or
broke the science, shifted a proposed experiment into the past tense, weakened a
methodological requirement, or altered a quoted reference title.

Two of these would be serious on their own:

  * "Caffeine ... is unlikely to be an informative tool" reverses the sentence
    that justifies the entire paper.
  * The concluding sentence became "may promote the adoption of the adenosinergic
    braking system", which does not mean anything.

Replacement preserves run formatting so subscripts such as A1R survive.

    python fix_rubriq_v4.py
"""
import io
import os

from docx import Document

SRC = r"C:\Users\Leon\Downloads\Epilepsy\Epilepsia_Open_Hypothesis_Manuscript_v4.docx"

FIXES = [
 # ---------------------------------------------------- reversals of meaning
 ("Caffeine, an adenosine receptor antagonist, is unlikely to be an informative "
  "tool.",
  "Caffeine, an adenosine receptor antagonist, is an unlikely but informative "
  "tool.",
  "REVERSAL: caffeine described as uninformative, which contradicts the paper"),

 ("Chronic caffeine exposure may promote the adoption of the adenosinergic "
  "braking system",
  "Chronic caffeine exposure may adapt the adenosinergic braking system",
  "REVERSAL: 'promote the adoption of' is not meaningful; the claim is adaptation"),

 ("Chronic caffeine may adapt to this brake",
  "Chronic caffeine may adapt this brake",
  "'adapt to' reverses which thing is adapted"),

 ("We propose a test that separates adaptation from drug withdrawal and residual",
  "We propose a test that separates adaptation from withdrawal and residual drug",
  "'residual' left dangling; the contrast is with residual drug"),

 # ------------------------------------------------------- factual precision
 ("the affinity of rodents for A\u2081R lies in the low tens of micromolar range",
  "rodent affinity for A\u2081R lies in the low tens of micromolar range",
  "it is caffeine's affinity for the rodent receptor, not rodents' affinity"),

 ("although actual exposure to drinking water depends on the intake pattern",
  "although actual exposure under drinking-water administration depends on the "
  "intake pattern",
  "exposure is to caffeine delivered in drinking water, not to water"),

 ("The highest dose is approximately 2.4 mg/kg/day in humans by body-surface "
  "area (14)",
  "The highest dose scales to about 2.4 mg/kg/day in humans by body-surface "
  "area (14)",
  "2.4 mg/kg/day is the human-equivalent of the 30 mg/kg/day mouse dose, not "
  "the dose given"),

 ("decreased susceptibility to bicuculline- and pentylenetetrazol (PTZ)-induced "
  "seizures (7) in one study in which no change in A\u2081R number was detected "
  "(8) and evidence of A\u2082A involvement was detected (9).",
  "decreased susceptibility to bicuculline- and pentylenetetrazol (PTZ)-induced "
  "seizures (7), in one study without any change in A\u2081R number (8), and "
  "with evidence for A\u2082A involvement (9).",
  "references 7, 8 and 9 are separate studies; the edit merged them into one"),

 # ------------------------------- a proposed experiment, not a completed one
 ("The mice received plain water (G1) or caffeine in the drinking water, which "
  "was increased from 5 to 30 mg/kg/day over 8 weeks and held to week 12, under "
  "continuous video-EEG with prespecified rules that step a mouse back if "
  "interictal spikes, electrographic seizures or behavioral overstimulation "
  "appeared.",
  "Mice receive plain water (G1) or caffeine in drinking water escalated from 5 "
  "to 30 mg/kg/day over 8 weeks and held to week 12, under continuous video-EEG "
  "with prespecified rules that step a mouse back if interictal spikes, "
  "electrographic seizures or behavioral overstimulation appear.",
  "past tense implies the experiment was performed; it is proposed"),

 ("was used to explore the delivery mode, washout schedule and sample size "
  "before any animals were used (19)",
  "was used to explore the delivery mode, washout schedule and sample size "
  "before any animals are used (19)",
  "the point is that no animals have been used yet"),

 ("C1. The PTZ threshold in G2 did not differ from that in G1 at 48 h, while "
  "residual caffeine and metabolites were confirmed to be negligible.",
  "C1. The PTZ threshold in G2 does not differ from that in G1 at 48 h while "
  "residual caffeine and metabolites are confirmed negligible.",
  "falsification criteria describe future results, not past ones"),

 ("If Hypothesis A was supported", "If Hypothesis A were supported",
  "subjunctive, and inconsistent with 'If C were supported' two clauses later"),
 ("If B was supported", "If B were supported", "same"),

 # -------------------------------------------- weakened requirements and logic
 ("No animal work is proposed here without prior institutional approval or "
  "registration.",
  "No animal work is proposed here without prior institutional approval and "
  "registration.",
  "'or' makes either sufficient; both are required"),

 ("A positive result is interpretable only if residual caffeine and metabolites "
  "are negligible at the time of testing.",
  "A positive result is interpretable only if residual caffeine and metabolites "
  "are shown to be negligible at the time of testing.",
  "they must be demonstrated negligible, not merely be negligible"),

 ("The mechanism may be real, and the experiment is still uninformative, which "
  "is a failure of the design",
  "The mechanism may be real and the experiment still uninformative, which is a "
  "failure of the design",
  "the concessive reading is the point: both can hold at once"),

 ("C8. Adaptation is measurable but reversed faster than approximately two "
  "days, with the design placed outside the powered region shown in Figure 4A.",
  "C8. Adaptation is measurable but reverses faster than approximately two "
  "days, placing the design outside the powered region shown in Figure 4A.",
  "tense, and the design is placed outside by the reversal"),

 # --------------------------------------------------------- garbled captions
 ("Table 2. Each of the core groups and the mechanistic contrast are provided.",
  "Table 2. Core groups and the mechanistic contrast each provides.",
  "caption garbled by the edit"),

 ("(G1 vehicle; G2 increased caffeine concentration and then withdrawn; G3 "
  "increased caffeine concentration and continued)",
  "(G1 vehicle; G2 escalated caffeine then withdrawn; G3 escalated caffeine "
  "continued)",
  "the protocol escalates dose, not concentration"),

 ("1  unresolved question", "1  The unresolved question", "heading lost its article"),
 ("Falsification criteria The hypothesis and this design are refuted",
  "Falsification criteria. The hypothesis and this design are refuted",
  "missing full stop after the run-in heading"),

 # ------------------------------------------------------------ ethics wording
 ("No animals were used in this hypothesis or in silico study.",
  "No animals were used in this hypothesis and in-silico study.",
  "it is one study that is both; 'or' reads as two alternatives"),

 # ------------------------------------- reference titles are quoted verbatim
 ("Caffeine-induced behavioral stimulation is dose- and concentration dependent.",
  "Caffeine-induced behavioural stimulation is dose- and concentration-dependent.",
  "published title is British-spelled and hyphenated (Br J Pharmacol 1990)"),
 ("chronic coadministration and acute withdrawal of caffeine and ethanol",
  "chronic co-administration and acute withdrawal of caffeine and ethanol",
  "published title is hyphenated (J Basic Clin Physiol Pharmacol 2017)"),
]


def replace_in_paragraph(p, old, new):
    runs = p.runs
    if not runs:
        return False
    text = "".join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    owner = []
    for k, r in enumerate(runs):
        owner.extend([k] * len(r.text))
    j = i + len(old)
    nt = text[:i] + new + text[j:]
    no = owner[:i] + [owner[i]] * len(new) + owner[j:]
    for k, r in enumerate(runs):
        r.text = "".join(c for c, o in zip(nt, no) if o == k)
    return True


def blocks(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p


doc = Document(SRC)
applied, missed = [], []
for old, new, why in FIXES:
    hit = False
    for p in blocks(doc):
        if replace_in_paragraph(p, old, new):
            hit = True
            break
    (applied if hit else missed).append((old, why))
doc.save(SRC)

log = ["APPLIED %d of %d" % (len(applied), len(FIXES)), ""]
for old, why in applied:
    log.append("  + %-56s  %s" % (old[:56], why))
if missed:
    log += ["", "NOT FOUND:"]
    for old, why in missed:
        log.append("  ! %-56s  %s" % (old[:56], why))
io.open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "_fix_v4_log.txt"),
        "w", encoding="utf-8").write("\n".join(log))
print("applied %d of %d; see _fix_v4_log.txt" % (len(applied), len(FIXES)))
if missed:
    print("MISSED %d" % len(missed))
